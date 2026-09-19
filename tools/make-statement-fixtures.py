#!/opt/homebrew/bin/python3
"""Builds the SYNTHETIC statement fixtures under tools/fixtures/statements/.

Every name, number and amount here is invented. The LAYOUTS are copied from three
real statement files (docs/specs/statement-capture-spec.md section 15), because the
layout is what the parser, the column mapper and the arithmetic proof have to
survive: a table header on row 25 under merged cells, sub-header rows inside the
table, totals rows after it, credits printed as negative numbers, newest-first row
order, text dates, failed rows that still carry an amount, and wallet rows that do
not move the wallet balance.

Real statements never enter this repo (research/statements/ is git-ignored).

    /opt/homebrew/bin/python3 tools/make-statement-fixtures.py

Needs openpyxl and msoffcrypto-tool. Deterministic: same bytes in, same rows out
(the encrypted file's salt differs per run, its contents do not).
"""
import io, os, json
from datetime import datetime, timedelta
import openpyxl
from msoffcrypto.format.ooxml import OOXMLFile

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'fixtures', 'statements')
os.makedirs(OUT, exist_ok=True)
HOLDER = 'TRAN THI MAI'
HOLDER_VI = 'Trần Thị Mai'
PASSWORD = '01011990'


def grid_of(wb):
    """The cell grid as the app's own .xlsx reader hands it over (42-xlsx-parse.js):
    one array per non-empty row, every cell a string, sparse rows left-aligned by
    column index. tools/statement-table.test.js reads this instead of re-implementing
    a ZIP + XML reader under Node."""
    ws = wb.worksheets[0]; out = []
    for row in ws.iter_rows():
        cells = ['' if c.value is None else (repr(c.value) if isinstance(c.value, float) else str(c.value)) for c in row]
        while cells and cells[-1] == '': cells.pop()
        if any(x != '' for x in cells): out.append(cells)
    return out


def save(wb, name, password=None):
    path = os.path.join(OUT, name)
    with open(os.path.join(OUT, name.replace('.locked', '').replace('.xlsx', '.grid.json')), 'w') as fh:
        json.dump(grid_of(wb), fh, ensure_ascii=False)
    if not password:
        wb.save(path)
        return
    plain = io.BytesIO()
    wb.save(plain)
    plain.seek(0)
    with open(path, 'wb') as fh:
        OOXMLFile(plain).encrypt(password, fh)


# --- 1. Bank account statement: preamble, header on row 15 in column B, -------
#        debit/credit columns, running balance, NEWEST FIRST, day-only text dates.
def bank_account():
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Sao_ke_tai_khoan'
    opening = 12_500_000
    moves = [  # oldest first: (day, description, debit, credit)
        (1, 'TRAN THI MAI chuyen tien', 350_000, 0),
        (2, 'Nap tien vao vi dien tu - MOMO 0900000001', 500_000, 0),
        (3, 'LUONG THANG 08 CONG TY ABC', 0, 18_000_000),
        (5, '512345xxxxxx6789-000000000123456 - TRAN THI MAI - Thanh toan sao ke the Master Card 08', 2_150_000, 0),
        (7, 'TRAN THI MAI chuyen tien den LE THI HOA - 0123456789', 1_200_000, 0),
        (9, 'FC12-0000012345', 45_000, 0),
        (12, 'Nap tien vao vi dien tu - MOMO 0900000001', 300_000, 0),
        (15, 'NGUYEN VAN BINH chuyen tien an trua', 0, 120_000),
        (18, 'TRAN THI MAI chuyen tien', 80_000, 0),
        (21, 'Phi duy tri tai khoan', 11_000, 0),
        (25, 'TRAN THI MAI chuyen tien den TRAN THI MAI - 9988776655', 2_000_000, 0),
        (28, 'Tra lai tien gui', 0, 3_250),
    ]
    bal = opening; rows = []
    for d, desc, db, cr in moves:
        bal = bal - db + cr
        rows.append(('%02d/08/2026' % d, desc, db or None, cr or None, bal))
    tdb = sum(m[2] for m in moves); tcr = sum(m[3] for m in moves)
    ws['B2'] = 'GIAO DỊCH TÀI KHOẢN'; ws['B3'] = 'ACCOUNT STATEMENT'
    pre = [('Số tài khoản (Account number)', ': 000111222333444'), ('Loại tiền (Currency)', ': VND'),
           ('Chủ tài khoản (Account holder)', ': ' + HOLDER), ('Số CCCD (ID number)', ': 000000000000'),
           ('Địa chỉ (Address)', ': 1 Duong So 1, Phuong 1, Quan 1, TP HCM'),
           ('Ngày in sao kê (Statement date)', ': 01/09/2026'),
           ('Thời gian sao kê (Period covered)', ': 01/08/2026 - 31/08/2026')]
    for i, (k, v) in enumerate(pre):
        ws.cell(5 + i, 2, k); ws.cell(5 + i, 3, v)
    ws['B12'] = 'Phát sinh nợ (Total debit)'; ws['C12'] = ': {:,} VND'.format(tdb)
    ws['D12'] = 'Số dư đầu kỳ (Opening balance)'; ws['E12'] = ': {:,} VND'.format(opening)
    ws['B13'] = 'Phát sinh có (Total credit)'; ws['C13'] = ': {:,} VND'.format(tcr)
    ws['D13'] = 'Số dư cuối kỳ (Closing balance)'; ws['E13'] = ': {:,} VND'.format(bal)
    for j, h in enumerate(['Ngày giao dịch/Transaction date', 'Nội dung/Description', 'Ghi nợ/Debit',
                           'Ghi có/Credit', 'Số dư cuối/Running balance']):
        ws.cell(15, 2 + j, h)
    for i, r in enumerate(reversed(rows)):                       # newest first
        for j, v in enumerate(r):
            if v is not None: ws.cell(16 + i, 2 + j, v)
    save(wb, 'bank-account.xlsx')
    return {'file': 'bank-account.xlsx', 'rows': len(rows), 'opening': opening, 'closing': bal,
            'total_debit': tdb, 'total_credit': tcr, 'order': 'newest-first'}


# --- 2. Credit card statement: LOCKED, 24-row merged preamble, header row 25, ---
#        sub-header rows inside the table, totals after it, credits NEGATIVE.
def credit_card():
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'saoke'
    ws['A1'] = 'Sao kê giao dịch thẻ tín dụng '; ws['A2'] = 'Credit Card Monthly Statement'
    ws['A3'] = 'BANK CASH BACK'; ws['A4'] = HOLDER
    prev_debt = 2_150_000
    pay = [('05/08/2026', '05/08/2026', '512345xxxxxx6789-000000000123456 - TRAN THI MAI - Thanh toan sao ke the Master Card 08',
            '6012-Member Financial Institution', 0.0, -2_150_000.0)]
    buys = [('02/08/2026', '04/08/2026', 'Mua Hàng / Foody                  ', '5812-Eating Places', 185_000.0, 0.0),
            ('06/08/2026', '07/08/2026', 'Mua Hàng / APPLE.COM/BILL', '5818-Digital Goods', 79_000.0, 0.0),
            ('10/08/2026', '12/08/2026', 'Mua Hàng / GRAB*A-1234', '4121-Taxicabs', 64_000.0, 0.0),
            ('14/08/2026', '15/08/2026', 'Mua Hàng / OPENAI *CHATGPT', '5734-Computer Software', 527_400.0, 0.0),
            ('14/08/2026', '15/08/2026', 'Phí Giao Dịch Ngoại Tệ', '5734-Computer Software', 15_822.0, 0.0),
            ('20/08/2026', '20/08/2026', 'Mua Hàng / WINMART Q1', '5411-Grocery Stores', 412_300.0, 0.0),
            ('27/08/2026', '28/08/2026', 'Mua Hàng / SHOPEE', '5311-Department Stores', 236_000.0, 0.0)]
    tdb = sum(b[4] for b in buys); tcr = sum(-p[5] for p in pay); end = prev_debt + tdb - tcr
    meta = [(5, 'Số thẻ chính (Primary Card Number)', '512345******6789'), (6, 'Số tài khoản thẻ (Card Account)', 'C000000000123456'),
            (7, 'Hạn mức tín dụng (Credit Limit)', 50_000_000.0), (8, 'Lãi suất mua sắm/ rút tiền (Purchase/Cash rate)', '2.75%/tháng'),
            (15, 'Tóm tắt sao kê (Statement summary)', None), (16, 'Ngày sao kê (Statement Date)', '31/08/2026'),
            (17, 'Ngày đến hạn thanh toán (Payment Due Date)', '15/09/2026'), (18, 'Phương thức thanh toán (Payment method)', 'Thu Nợ Tối Thiểu'),
            (19, 'Số TK trích nợ (Auto Debit Account)', '000111222333444'), (20, 'Dư nợ kỳ trước (VND) (Previous Balance)', float(prev_debt)),
            (21, 'Phát sinh nợ trong kỳ (VND) (Total Debit)', tdb), (22, 'Phát sinh có trong kỳ (VND) (Total Credit)', tcr),
            (23, 'Dư nợ cuối kỳ (VND) (End Balance)', end), (24, 'Thanh toán tối thiểu (VND) (Minimum payment)', round(end * 0.05))]
    for r, k, v in meta:
        ws.cell(r, 1, k)
        if v is not None: ws.cell(r, 3, v)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    for r in (9, 10, 11, 12, 13, 14):
        ws.cell(r, 1, 'Tên chủ thẻ phụ (Sup. Card holder)' if r % 2 else 'Số thẻ phụ (Sup. Card Number)')
    hdr = ['Ngày giao dịch\nTransaction date', 'Ngày hạch toán\nPost date', 'Diễn giải\nDetails', 'MCC\nMCC',
           'Ghi nợ/Debit\n(VND)', 'Ghi có/Credit\n(VND)']
    for j, h in enumerate(hdr): ws.cell(25, 1 + j, h)
    r = 26
    ws.cell(r, 1, 'Số thẻ/ Số tài khoản\nCard number / Account'); ws.cell(r, 3, 'C000000000123456'); r += 1
    for row in pay:
        for j, v in enumerate(row): ws.cell(r, 1 + j, v)
        r += 1
    ws.cell(r, 1, 'Số thẻ/ Số tài khoản\nCard number / Account'); ws.cell(r, 3, '512345******6789'); r += 1
    for row in buys:
        for j, v in enumerate(row): ws.cell(r, 1 + j, v)
        r += 1
    for k, v in [('Phát sinh nợ trong kỳ (VND) (Total Debit)', tdb), ('Phát sinh có trong kỳ (VND) (Total Credit)', tcr),
                 ('Dư nợ kỳ trước (VND) (Previous Balance)', float(prev_debt)), ('Dư nợ cuối kỳ (VND) (End Balance)', end),
                 ('Thanh toán tối thiểu (VND) (Minimum payment)', round(end * 0.05))]:
        ws.cell(r, 1, k); ws.cell(r, 5, v); r += 1
    save(wb, 'credit-card.locked.xlsx', PASSWORD)
    return {'file': 'credit-card.locked.xlsx', 'password': PASSWORD, 'rows': len(pay) + len(buys),
            'previous_debt': prev_debt, 'closing_debt': end, 'total_debit': tdb, 'total_credit': tcr,
            'credit_sign': 'negative'}


# --- 3. E-wallet statement: LOCKED, headers on row 1, one signed amount column, --
#        to-the-second text timestamps, NEWEST FIRST, failed rows, and debit rows
#        that do NOT move the wallet balance (paid from a linked bank).
def ewallet():
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'LSGD'
    hdr = ['STT', 'Thời gian ', 'Mã giao dịch', 'Loại giao dịch', 'Tài khoản chuyển', 'Tên định danh\n Tài khoản chuyển',
           'Tài khoản nhận', 'Tên định danh \nTài khoản nhận', 'Số Tiền', 'Số Dư Sau giao dịch', 'Trạng Thái GD']
    for j, h in enumerate(hdr): ws.cell(1, 1 + j, h)
    W = '0900000001'; t0 = datetime(2026, 8, 1, 8, 15, 0); bal = 240_000.0
    ev = [  # oldest first: (type, from_acct, from_name, to_acct, to_name, amount, moves_balance, ok)
        ('GRAB', W, 'Ví MoMo', 'm4becomgrab_moca_v2', 'GRAB', -42_000, True, True),
        ('Nạp tiền vào Ví để thanh toán GRAB', 'vcb01.02.bank', 'Vietcombank', W, 'GRAB', 500_000, True, True),
        ('Thanh toán EVERY HALF COFFEE ROASTERS', W, HOLDER, 'm4bopceveryhalfcoffee', 'EVERY HALF COFFEE ROASTERS', -65_000, True, True),
        ('Nhận tiền từ Nguyễn Văn Bình', '*******123', 'Nguyễn Văn Bình', W, HOLDER_VI, 120_000, True, True),
        ('GRAB', W, 'Ví MoMo', 'm4becomgrab_moca_v2', 'GRAB', -38_000, True, False),                 # failed
        ('Thanh toán SUKIYA', W, 'Ngân hàng liên kết', 'm4b_sukiya', 'SUKIYA', -74_000, False, True),   # bank-funded
        ('Chuyển đến Lê Thị Hoa', W, HOLDER_VI, '*******456', 'Lê Thị Hoa', -300_000, True, True),
        ('Nhận lương từ CONG TY ABC', 'accounting_abc', 'CONG TY ABC', W, HOLDER, 2_000_000, True, True),
        ('Hoàn tiền giao dịch từ Đối tác', 'm4becomgrab_moca_v2', 'Đối tác MoMo', W, 'Đối tác MoMo', 12_000, True, True),
        ('Thanh toán cho Dong Tay Barber', W, 'VIB', 'm4b_dongtay', 'Dong Tay Barber', -90_000, False, True),  # bank-funded
        ('Nạp tiền điện thoại Mobifone', W, HOLDER, 'm4b_mbp_airtime', 'Mobifone', -50_000, True, True),
        ('Chuyển tiền qua mã QR đến Phạm Quốc Huy', W, HOLDER_VI, '*******789', 'Phạm Quốc Huy', -45_000, True, True),
        ('Nạp tiền vào Ví để thanh toán', 'vcb01.02.bank', 'Vietcombank', W, HOLDER, 300_000, True, True),
        ('Thanh toán REVI COFFEE', W, HOLDER, 'm4b_topbrand12_revi', 'REVI COFFEE', -55_000, True, True),
    ]
    rows = []
    for i, (typ, fa, fn, ta, tn, amt, moves, ok) in enumerate(ev):
        ts = t0 + timedelta(days=i * 2, minutes=i * 7, seconds=i * 3)
        if ok and moves: bal += amt
        rows.append([None, ts.strftime('%d/%m/%Y %H:%M:%S'), str(90000000000 + i * 7919), typ, fa, fn, ta, tn,
                     float(amt), float(bal), 'Thành công' if ok else 'Thất bại'])
    rows.reverse()                                                # newest first
    for i, r in enumerate(rows):
        r[0] = float(i + 1)
        for j, v in enumerate(r): ws.cell(2 + i, 1 + j, v)
    save(wb, 'ewallet.locked.xlsx', PASSWORD)
    return {'file': 'ewallet.locked.xlsx', 'password': PASSWORD, 'rows': len(rows), 'failed': 1,
            'bank_funded': 2, 'order': 'newest-first', 'closing': bal}


manifest = {'note': 'SYNTHETIC. Layouts copied from real statements; every value invented.',
            'fixtures': [bank_account(), credit_card(), ewallet()]}
with open(os.path.join(OUT, 'manifest.json'), 'w') as fh:
    json.dump(manifest, fh, ensure_ascii=False, indent=2)
print(json.dumps(manifest, ensure_ascii=False, indent=2))
